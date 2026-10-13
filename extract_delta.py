#!/usr/bin/env python3
"""
Keep only the yellow-highlighted rows' blocks from TI+ delta XMLs.

Usage:  python extract_delta.py DeltaReport.xlsx deltas.zip out_dir
Needs:  pip install openpyxl lxml
"""
import re
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook

CTRL_NS = "urn:control.services.tiplus2.misys.com"
M_NS = "urn:messages.service.ti.apps.tiplus2.misys.com"

DEFAULT_KEY = ["ParameterSet", "Area", "Category", "Name"]
# Per-operation key overrides: Excel header name == XML child tag name
KEYS = {
    # "Branch": ["BranchCode"],
}
ACTION_TO_MAINT = {"DELETED": "D", "UPDATED": "U", "ADDED": "A"}


def is_yellow(cell):
    f = cell.fill
    if not f or f.fill_type != "solid":
        return False
    c = f.fgColor
    if c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8:
        r, g, b = (int(c.rgb[i:i + 2], 16) for i in (2, 4, 6))
        return r >= 0xE6 and g >= 0xD0 and b <= 0xB0
    if c.type == "indexed":
        return c.indexed in (5, 13, 43, 44)
    return False  # theme colours: inspect and add handling if your yellow is a theme colour


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def find_header(ws):
    for r in range(1, 16):
        vals = [norm(c.value) for c in ws[r]]
        if "Action" in vals:
            return r, {v: i for i, v in enumerate(vals) if v}
    return None, {}


def find_xml(files, op):
    suffix = f".{op}-delta".lower()
    for stem, p in files.items():
        if stem.endswith(suffix):
            return p
    for stem, p in files.items():  # sheet names are capped at 31 chars -> prefix fallback
        if f".{op}".lower() in stem:
            return p
    return None


def highlighted_keys(ws, key_fields):
    hr, cols = find_header(ws)
    if hr is None:
        return None, []
    missing = [k for k in key_fields if k not in cols]
    if missing:
        raise SystemExit(f"[{ws.title}] key columns not in sheet: {missing} (set KEYS)")
    keys = []
    for row in ws.iter_rows(min_row=hr + 1):
        if not any(is_yellow(c) for c in row):
            continue
        key = tuple(norm(row[cols[k]].value) for k in key_fields)
        maint = ACTION_TO_MAINT.get(norm(row[cols["Action"]].value).upper(), "")
        keys.append((key, maint))
    return hr, keys


def process(ws, files, out_dir):
    op = re.sub(r"-\d+$", "", ws.title)
    key_fields = KEYS.get(op, DEFAULT_KEY)
    hr, wanted = highlighted_keys(ws, key_fields)
    if not wanted:
        return
    src = find_xml(files, op)
    if not src:
        print(f"[{ws.title}] no XML found in zip for '{op}'")
        return

    parser = etree.XMLParser(remove_blank_text=True, huge_tree=True)
    tree = etree.parse(str(src), parser)
    item = tree.getroot().find(f"{{{M_NS}}}ItemRequest")
    blocks = item.findall(f"{{{CTRL_NS}}}ServiceRequest")

    for c in item.xpath("comment()"):  # drop the template's row comments
        item.remove(c)

    want = set(wanted)
    found = set()
    for b in blocks:
        rec = next((e for e in b if isinstance(e.tag, str) and e.tag.startswith(f"{{{M_NS}}}")), None)
        if rec is None:
            item.remove(b)
            continue
        key = tuple((rec.findtext(f"{{{M_NS}}}{k}") or "").strip() for k in key_fields)
        maint = (rec.findtext(f"{{{M_NS}}}MaintType") or "").strip()
        hit = (key, maint) if (key, maint) in want else (key, "") if (key, "") in want else None
        if hit:
            found.add(hit)
        else:
            item.remove(b)

    out = Path(out_dir) / src.name
    tree.write(str(out), xml_declaration=True, encoding="utf-8", pretty_print=True)
    print(f"[{ws.title}] highlighted={len(want)} kept={len(found)} of {len(blocks)} -> {out.name}")
    for k in want - found:
        print(f"    NOT FOUND in XML: {k}")


def main(xlsx, zip_path, out_dir):
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    tmp = tempfile.mkdtemp()
    with zipfile.ZipFile(zip_path) as z:
        z.extractall(tmp)
    files = {p.stem.lower(): p for p in Path(tmp).rglob("*") if p.is_file()}
    files = {k.lower(): v for k, v in files.items()}

    wb = load_workbook(xlsx)  # full mode: fills need styles
    for ws in wb.worksheets:
        if ws.title.lower() == "summary":
            continue
        process(ws, {k: v for k, v in files.items()}, out_dir)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(*sys.argv[1:4])
